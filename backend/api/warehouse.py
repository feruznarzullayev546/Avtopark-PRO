import json
from datetime import date
from django.db import transaction
from django.db.models import Sum
from django.http import JsonResponse, HttpResponse
from django.views.decorators.csrf import csrf_protect
from django.views.decorators.http import require_http_methods
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from core.models import Part, Receipt, Issue, Vehicle, DriverProfile, MaintenanceSchedule, UserProfile, WarehouseReturn, WarehouseAdjustment, OtherExpense, AuditLog
from .trips import _company, _profile, _can_manage, _errors
from .branch_scope import branch_kwargs, relation_branch_kwargs


def _warehouse_role(request):
    if request.user.is_superuser or UserProfile.objects.filter(user=request.user, role=UserProfile.ROLE_SUPER_ADMIN).exists():
        return "SUPER_ADMIN"
    c = _company(request)
    if not c:
        return None
    m = request.user.company_memberships.filter(company=c, status="ACTIVE").first()
    return m.role if m else None

def _warehouse_read(request):
    return _warehouse_role(request) in {"SUPER_ADMIN", UserProfile.ROLE_DIRECTOR, UserProfile.ROLE_MANAGER, UserProfile.ROLE_WAREHOUSE, UserProfile.ROLE_ACCOUNTANT, UserProfile.ROLE_DISPATCHER}

def _warehouse_create(request):
    return _warehouse_role(request) in {"SUPER_ADMIN", UserProfile.ROLE_DIRECTOR, UserProfile.ROLE_WAREHOUSE}

def _warehouse_update(request):
    return _warehouse_role(request) in {"SUPER_ADMIN", UserProfile.ROLE_DIRECTOR, UserProfile.ROLE_MANAGER, UserProfile.ROLE_WAREHOUSE}

def _warehouse_direct(request):
    return _warehouse_role(request) in {"SUPER_ADMIN", UserProfile.ROLE_DIRECTOR}

def _audit(request, company, action, obj, description, metadata=None):
    AuditLog.objects.create(
        company=company, user=request.user, action=action, model_name=obj.__class__.__name__,
        object_id=str(getattr(obj, "pk", "")), description=description,
        branch=getattr(obj, "branch", None), metadata=metadata or {},
        ip_address=request.META.get("REMOTE_ADDR"), user_agent=request.META.get("HTTP_USER_AGENT", "")[:500]
    )

def _json(request):
    try:
        return json.loads(request.body or '{}')
    except json.JSONDecodeError:
        return None


def _date(value):
    if not value:
        return date.today()
    try:
        return date.fromisoformat(str(value))
    except (TypeError, ValueError):
        raise ValueError("Sana noto'g'ri")


def _num(value, name, minimum=None, required=True):
    if value in (None, ''):
        if required:
            raise ValueError(f"{name} majburiy")
        return None
    try:
        n = float(value)
    except (TypeError, ValueError):
        raise ValueError(f"{name} noto'g'ri")
    if minimum is not None and n < minimum:
        raise ValueError(f"{name} {minimum:g} dan kichik bo'lishi mumkin emas")
    return n


def _weighted_unit_price(part):
    from django.db.models import Sum
    incoming = float(Receipt.objects.filter(part=part, branch_id=part.branch_id).aggregate(v=Sum('quantity'))['v'] or 0)
    incoming_value = float(Receipt.objects.filter(part=part, branch_id=part.branch_id).aggregate(v=Sum('quantity__mul_unit_price'))['v'] or 0) if False else 0
    # Django cannot portablely multiply fields without ExpressionWrapper in older versions.
    incoming_value = sum(float(r.quantity) * float(r.unit_price) for r in Receipt.objects.filter(part=part, branch_id=part.branch_id).only('quantity','unit_price'))
    returned = list(WarehouseReturn.objects.filter(part=part, branch_id=part.branch_id).only('quantity','unit_price'))
    ret_qty = sum(float(x.quantity) for x in returned)
    ret_value = sum(float(x.quantity) * float(x.unit_price) for x in returned)
    if incoming + ret_qty <= 0:
        return float(part.unit_price or 0)
    return (incoming_value + ret_value) / (incoming + ret_qty)

def _stock_for_part(part):
    incoming = float(part.receipts.filter(branch_id=part.branch_id).aggregate(v=Sum('quantity'))['v'] or 0)
    returned = float(part.returns.filter(branch_id=part.branch_id).aggregate(v=Sum('quantity'))['v'] or 0)
    issues = float(part.issues.filter(branch_id=part.branch_id).aggregate(v=Sum('quantity'))['v'] or 0)
    adjustments = float(part.adjustments.filter(branch_id=part.branch_id).aggregate(v=Sum('quantity'))['v'] or 0)
    return incoming + returned + adjustments - issues

def _part_json(p):
    stock = max(_stock_for_part(p), 0)
    avg_price = _weighted_unit_price(p)
    return {
        'id': p.id, 'code': p.code, 'name': p.name, 'category': p.category,
        'model': p.model, 'unit': p.unit, 'min_stock': float(p.min_stock or 0),
        'unit_price': avg_price, 'stock': stock,
        'stock_value': stock * avg_price,
        'status': 'CRITICAL' if stock <= 0 else ('LOW' if stock <= float(p.min_stock or 0) else 'OK'),
    }

def _receipt_json(r):
    return {
        'id': r.id, 'date': r.date.isoformat(), 'part_id': r.part_id,
        'part': _part_json(r.part), 'compatible_model': r.compatible_model,
        'quantity': float(r.quantity), 'unit_price': float(r.unit_price),
        'total': float(r.total), 'supplier': r.supplier,
        'document_no': r.document_no, 'note': r.note,
    }


def _issue_json(i):
    return {
        'id': i.id, 'date': i.date.isoformat(), 'part_id': i.part_id,
        'part': _part_json(i.part), 'vehicle_id': i.vehicle_id,
        'vehicle': i.vehicle.number if i.vehicle_id else '',
        'driver_id': i.driver_id,
        'driver': i.historical_driver_name,
        'maintenance_id': i.maintenance_id,
        'maintenance': i.maintenance.name if i.maintenance_id else '',
        'odometer': float(i.odometer or 0), 'quantity': float(i.quantity),
        'unit_price': float(i.unit_price), 'total': float(i.total),
        'service_type': i.service_type, 'note': i.note,
    }



def _return_json(r):
    return {
        'id': r.id, 'date': r.date.isoformat(), 'issue_id': r.issue_id,
        'part_id': r.part_id, 'part': _part_json(r.part), 'vehicle_id': r.vehicle_id,
        'vehicle': r.vehicle.number, 'quantity': float(r.quantity),
        'unit_price': float(r.unit_price), 'total': float(r.total),
        'document_no': r.document_no, 'reason': r.reason, 'note': r.note,
    }

def _adjustment_json(a):
    return {
        'id': a.id, 'date': a.date.isoformat(), 'part_id': a.part_id,
        'part': _part_json(a.part), 'quantity': float(a.quantity),
        'unit_price': float(a.unit_price), 'total': float(a.total),
        'reason': a.reason, 'document_no': a.document_no, 'note': a.note,
        'type': 'SURPLUS' if a.quantity > 0 else 'SHORTAGE',
    }

def _meta(request, c):
    vehicles = Vehicle.objects.filter(company=c, **branch_kwargs(request), active=True).order_by('number')
    drivers = DriverProfile.objects.filter(company=c, **branch_kwargs(request), active=True).select_related('vehicle').order_by('full_name')
    schedules = MaintenanceSchedule.objects.filter(
        vehicle__company=c, **relation_branch_kwargs(request, 'vehicle__branch'), active=True
    ).select_related('vehicle', 'part').order_by('vehicle__number', 'name')
    return {
        'vehicles': [{'id': v.id, 'number': v.number, 'model': v.model,
                      'mileage': float(v.current_mileage or 0)} for v in vehicles],
        'drivers': [{'id': d.id, 'name': d.full_name, 'vehicle_id': d.vehicle_id,
                     'vehicle': d.vehicle.number if d.vehicle_id else ''} for d in drivers],
        'maintenance': [{'id': s.id, 'name': s.name, 'vehicle_id': s.vehicle_id,
                         'vehicle': s.vehicle.number, 'part_id': s.part_id} for s in schedules],
    }


@require_http_methods(['GET'])
def warehouse(request):
    if not request.user.is_authenticated or not _can_manage(request):
        return JsonResponse({'error': 'Ruxsat yo‘q'}, status=403)
    c = _company(request)
    if not c:
        return JsonResponse({'parts': [], 'receipts': [], 'issues': [], 'summary': {}, 'meta': {}})
    parts = list(Part.objects.filter(company=c, **branch_kwargs(request)).order_by('name', 'code'))
    receipts = Receipt.objects.filter(company=c, **branch_kwargs(request)).select_related('part').order_by('-date', '-id')[:100]
    issues = Issue.objects.filter(company=c, **branch_kwargs(request)).select_related('part', 'vehicle', 'driver', 'maintenance').order_by('-date', '-id')[:100]
    returns = WarehouseReturn.objects.filter(company=c, **branch_kwargs(request)).select_related('part','vehicle','issue').order_by('-date','-id')[:100]
    adjustments = WarehouseAdjustment.objects.filter(company=c, **branch_kwargs(request)).select_related('part').order_by('-date','-id')[:100]
    part_json = [_part_json(p) for p in parts]
    return JsonResponse({
        'parts': part_json,
        'receipts': [_receipt_json(x) for x in receipts],
        'issues': [_issue_json(x) for x in issues],
        'returns': [_return_json(x) for x in returns],
        'adjustments': [_adjustment_json(x) for x in adjustments],
        'summary': {
            'items': len(parts),
            'quantity': sum(x['stock'] for x in part_json),
            'value': sum(x['stock_value'] for x in part_json),
            'low': sum(1 for x in part_json if x['status'] == 'LOW'),
            'receipts_total': sum(float(x.total) for x in receipts),
            'issues_total': sum(float(x.total) for x in issues),
            'returns_total': sum(float(x.total) for x in returns),
            'adjustments_total': sum(float(x.total) for x in adjustments),
        },
        'meta': _meta(request, c),
    })


@require_http_methods(['GET'])
def warehouse_meta(request):
    if not request.user.is_authenticated or not _can_manage(request):
        return JsonResponse({'error': 'Ruxsat yo‘q'}, status=403)
    c = _company(request)
    return JsonResponse({'parts': [_part_json(p) for p in Part.objects.filter(company=c, **branch_kwargs(request)).order_by('name', 'code')],
                         'meta': _meta(request, c)})


@csrf_protect
@require_http_methods(['POST'])
def part_create(request):
    if not request.user.is_authenticated or not _warehouse_create(request):
        return JsonResponse({'error': 'Ruxsat yo‘q'}, status=403)
    c = _company(request); data = _json(request)
    if data is None: return JsonResponse({'error': 'Noto‘g‘ri JSON'}, status=400)
    try:
        code = str(data.get('code', '')).strip()
        name = str(data.get('name', '')).strip()
        if not code or not name: raise ValueError('Kod va nom majburiy')
        min_stock = _num(data.get('min_stock', 0), 'Minimal qoldiq', 0)
        price = _num(data.get('unit_price', 0), 'Birlik narxi', 0)
        if Part.objects.filter(company=c, **branch_kwargs(request), code=code).exists():
            return JsonResponse({'error': 'Bu mahsulot kodi allaqachon mavjud'}, status=400)
        p = Part.objects.create(company=c, branch=getattr(request,"active_branch",None), code=code, name=name,
            category=str(data.get('category', '')).strip(), model=str(data.get('model', '')).strip(),
            unit=str(data.get('unit', 'dona')).strip() or 'dona',
            min_stock=min_stock, unit_price=price)
    except ValueError as e:
        return JsonResponse({'error': str(e)}, status=400)
    _audit(request, c, 'CREATE', p, f'Ombor mahsuloti yaratildi: {p.name}', {'code': p.code})
    return JsonResponse({'part': _part_json(p)}, status=201)


@csrf_protect
@require_http_methods(['PATCH'])
def part_update(request, pk):
    if not request.user.is_authenticated or not _warehouse_update(request):
        return JsonResponse({'error': 'Ruxsat yo‘q'}, status=403)
    c = _company(request); p = Part.objects.filter(company=c, **branch_kwargs(request), pk=pk).first()
    if not p: return JsonResponse({'error': 'Mahsulot topilmadi'}, status=404)
    data = _json(request) or {}
    try:
        if 'code' in data:
            code = str(data['code']).strip()
            if not code: raise ValueError('Kod majburiy')
            if Part.objects.filter(company=c, **branch_kwargs(request), code=code).exclude(pk=p.pk).exists():
                raise ValueError('Bu mahsulot kodi allaqachon mavjud')
            p.code = code
        if 'name' in data: p.name = str(data['name']).strip()
        if 'category' in data: p.category = str(data['category']).strip()
        if 'model' in data: p.model = str(data['model']).strip()
        if 'unit' in data: p.unit = str(data['unit']).strip() or 'dona'
        if 'min_stock' in data: p.min_stock = _num(data['min_stock'], 'Minimal qoldiq', 0)
        if 'unit_price' in data: p.unit_price = _num(data['unit_price'], 'Birlik narxi', 0)
        p.save()
    except ValueError as e:
        return JsonResponse({'error': str(e)}, status=400)
    return JsonResponse({'part': _part_json(p)})


def _get_part(request, c, value):
    try: return Part.objects.get(company=c, **branch_kwargs(request), pk=int(value))
    except (Part.DoesNotExist, TypeError, ValueError): return None


@csrf_protect
@require_http_methods(['POST'])
def receipt_create(request):
    return _receipt_save(request)


@csrf_protect
@require_http_methods(['PATCH'])
def receipt_update(request, pk):
    return _receipt_save(request, pk)


def _receipt_save(request, pk=None):
    if not request.user.is_authenticated or not (_warehouse_update(request) if pk else _warehouse_create(request)):
        return JsonResponse({'error': 'Ruxsat yo‘q'}, status=403)
    c = _company(request); data = _json(request) or {}
    instance = Receipt.objects.filter(company=c, **branch_kwargs(request), pk=pk).first() if pk else None
    if pk and not instance: return JsonResponse({'error': 'Kirim topilmadi'}, status=404)
    try:
        part = _get_part(request, c, data.get('part_id'))
        if not part: raise ValueError('Mahsulot topilmadi')
        qty = _num(data.get('quantity'), 'Miqdor', 0)
        if qty <= 0: raise ValueError('Miqdor 0 dan katta bo‘lishi kerak')
        price = _num(data.get('unit_price', part.unit_price), 'Birlik narxi', 0)
        dt = _date(data.get('date'))
        # Receipt itself is always positive. Lock the affected part so concurrent
        # receipt/issue writes cannot observe an inconsistent stock snapshot.
        with transaction.atomic():
            part = Part.objects.select_for_update().get(pk=part.pk, company=c)
            r = instance or Receipt(company=c)
            r.date, r.part, r.quantity, r.unit_price = dt, part, qty, price
            r.compatible_model = str(data.get('compatible_model', '')).strip()
            r.supplier = str(data.get('supplier', '')).strip()
            r.document_no = str(data.get('document_no', '')).strip()
            r.note = str(data.get('note', '')).strip()
            r.company = c; r.branch = getattr(request,"active_branch",None); r.save()
            if abs(float(part.unit_price) - price) > 1e-9:
                part.unit_price = price; part.save(update_fields=['unit_price'])
    except ValueError as e:
        return JsonResponse({'error': str(e)}, status=400)
    _audit(request, c, 'UPDATE', r, f'Ombor kirimi o‘zgartirildi: {r.part.name}', {'quantity': float(r.quantity), 'document_no': r.document_no})
    return JsonResponse({'receipt': _receipt_json(r)})


@csrf_protect
@require_http_methods(['POST'])
def issue_create(request):
    return _issue_save(request)


@csrf_protect
@require_http_methods(['PATCH'])
def issue_update(request, pk):
    return _issue_save(request, pk)


def _issue_save(request, pk=None):
    if not request.user.is_authenticated or not (_warehouse_update(request) if pk else _warehouse_create(request)):
        return JsonResponse({'error': 'Ruxsat yo‘q'}, status=403)
    c = _company(request); data = _json(request) or {}
    instance = Issue.objects.filter(company=c, **branch_kwargs(request), pk=pk).first() if pk else None
    if pk and not instance: return JsonResponse({'error': 'Chiqim topilmadi'}, status=404)
    try:
        part = _get_part(request, c, data.get('part_id'))
        if not part: raise ValueError('Mahsulot topilmadi')
        vehicle = Vehicle.objects.filter(company=c, **branch_kwargs(request), active=True, pk=data.get('vehicle_id')).first()
        if not vehicle: raise ValueError('Faol avtomobilni tanlang')
        driver = DriverProfile.objects.filter(company=c, **branch_kwargs(request), active=True, pk=data.get('driver_id')).first() if data.get('driver_id') else None
        maintenance = MaintenanceSchedule.objects.filter(vehicle__company=c, **relation_branch_kwargs(request, 'vehicle__branch'), active=True, pk=data.get('maintenance_id')).first() if data.get('maintenance_id') else None
        if driver and driver.vehicle_id and driver.vehicle_id != vehicle.id:
            raise ValueError('Haydovchi boshqa avtomobilga biriktirilgan')
        if maintenance and maintenance.vehicle_id != vehicle.id:
            raise ValueError('Servis rejasi tanlangan avtomobilga tegishli emas')
        qty = _num(data.get('quantity'), 'Miqdor', 0)
        if qty <= 0: raise ValueError('Miqdor 0 dan katta bo‘lishi kerak')
        price = _num(data.get('unit_price', part.unit_price), 'Birlik narxi', 0)
        dt = _date(data.get('date'))
        with transaction.atomic():
            old_part_id = instance.part_id if instance else None
            ids = {part.pk}
            if old_part_id: ids.add(old_part_id)
            locked = {x.pk: x for x in Part.objects.select_for_update().filter(company=c, **branch_kwargs(request), pk__in=ids)}
            target = locked.get(part.pk)
            if not target: raise ValueError('Mahsulot topilmadi')
            available = float(target.stock())
            if instance and old_part_id == part.pk:
                available += float(instance.quantity)
            if qty > available + 1e-9:
                raise ValueError(f'Qoldiq yetarli emas. Mavjud: {available:g} {target.unit}.')
            old_maintenance = instance.maintenance if instance and instance.maintenance_id else None
            i = instance or Issue(company=c)
            i.company=c; i.branch=getattr(request,"active_branch",None); i.date=dt; i.part=target; i.vehicle=vehicle; i.driver=driver
            i.maintenance=maintenance; i.quantity=qty; i.unit_price=price
            i.odometer=_num(data.get('odometer', vehicle.current_mileage), 'Probeg', 0, required=False) or 0
            i.service_type = maintenance.name if maintenance else str(data.get('service_type', '')).strip()
            i.note = str(data.get('note', '')).strip()
            i.part_name_snapshot=target.name; i.part_code_snapshot=target.code
            i.unit_snapshot=target.unit; i.driver_name_snapshot=driver.full_name if driver else ''
            i.save()
            # Keep warehouse issues synchronized with Finance as one source of truth.
            exp = i.expense
            if not exp:
                exp = OtherExpense.objects.create(company=c, date=dt, vehicle=vehicle,
                    category='Sklad', description=f'Sklad chiqimi: {target.name}',
                    amount=float(i.total), note=i.note)
                i.expense = exp
                i.save(update_fields=['expense'])
            else:
                returned_total=float(i.returns.aggregate(v=Sum('quantity'))['v'] or 0) * float(i.unit_price)
                exp.date=dt; exp.vehicle=vehicle; exp.amount=max(0,float(i.total)-returned_total)
                exp.description=f'Sklad chiqimi: {target.name}'; exp.note=i.note
                exp.save(update_fields=['date','vehicle','amount','description','note'])
            if old_maintenance and old_maintenance.pk != i.maintenance_id:
                from core.services import sync_maintenance_from_issues
                sync_maintenance_from_issues(old_maintenance)
            if i.maintenance_id:
                from core.services import sync_maintenance_from_issues
                sync_maintenance_from_issues(i.maintenance)
    except ValueError as e:
        return JsonResponse({'error': str(e)}, status=400)
    _audit(request, c, 'UPDATE', i, f'Ombor chiqimi o‘zgartirildi: {i.part.name}', {'quantity': float(i.quantity), 'vehicle_id': i.vehicle_id})
    return JsonResponse({'issue': _issue_json(i)})

@require_http_methods(['GET'])
def replacement_history(request):
    if not request.user.is_authenticated or not _can_manage(request):
        return JsonResponse({'error': 'Ruxsat yo\'q'}, status=403)
    c = _company(request)
    qs = Issue.objects.filter(company=c, **branch_kwargs(request)).select_related('part','vehicle','driver','maintenance').order_by('-date','-id')
    vehicle_id = request.GET.get('vehicle')
    service = (request.GET.get('service') or '').strip()
    if vehicle_id:
        qs = qs.filter(vehicle_id=vehicle_id)
    if service:
        from django.db.models import Q
        qs = qs.filter(Q(service_type__icontains=service) | Q(part__name__icontains=service) | Q(part__code__icontains=service) | Q(note__icontains=service))
    rows = list(qs[:500])
    data = [_issue_json(x) for x in rows]
    return JsonResponse({
        'results': data,
        'summary': {
            'count': len(data),
            'quantity': sum(float(x['quantity']) for x in data),
            'total': sum(float(x['total']) for x in data),
        },
        'meta': _meta(request, c),
    })


@require_http_methods(['GET'])
def replacement_export(request):
    if not request.user.is_authenticated or not _can_manage(request):
        return JsonResponse({'error': 'Ruxsat yo\'q'}, status=403)
    c = _company(request)
    qs = Issue.objects.filter(company=c, **branch_kwargs(request)).select_related('part','vehicle','driver','maintenance').order_by('-date','-id')
    vehicle_id = request.GET.get('vehicle')
    if vehicle_id:
        qs = qs.filter(vehicle_id=vehicle_id)
    rows = list(qs[:5000])
    wb = Workbook(); ws = wb.active; ws.title = 'Almashtirishlar'
    headers = ['№','Sana','Avtomobil','Model','Ehtiyot qism','Artikul','Miqdor','Birlik','Birlik narxi','Jami summa','Probeg','Haydovchi','Servis','Izoh']
    ws.append(headers)
    for cell in ws[1]: cell.font = Font(bold=True); cell.alignment = Alignment(horizontal='center', wrap_text=True)
    total = 0
    for i, x in enumerate(rows, 1):
        total += float(x.total)
        ws.append([i,x.date.strftime('%d.%m.%Y'),x.vehicle.number,x.vehicle.model,x.historical_part_name,x.historical_part_code,x.quantity,x.historical_unit,x.unit_price,x.total,x.odometer,x.historical_driver_name,x.service_type or (x.maintenance.name if x.maintenance_id else ''),x.note])
    ws.append([]); ws.append(['','','','','','','','','JAMI',total])
    ws.cell(ws.max_row,9).font = Font(bold=True); ws.cell(ws.max_row,10).font = Font(bold=True)
    widths=[6,13,16,22,30,16,12,12,16,18,14,24,24,32]
    for i,w in enumerate(widths,1): ws.column_dimensions[ws.cell(1,i).column_letter].width=w
    ws.freeze_panes='A2'; ws.auto_filter.ref=f'A1:N{max(1,ws.max_row-2)}'
    from io import BytesIO
    out=BytesIO(); wb.save(out); out.seek(0)
    return HttpResponse(out.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="avtopark_almashtirishlar_tarixi.xlsx"'})


@require_http_methods(['GET'])
def warehouse_export(request):
    if not request.user.is_authenticated or not _can_manage(request):
        return JsonResponse({'error':'Ruxsat yo‘q'},status=403)
    c=_company(request)
    wb=Workbook(); ws=wb.active; ws.title='Sklad qoldiq'
    headers=['Kod','Mahsulot','Kategoriya','Model','Birlik','Qoldiq','Min.','Weighted avg','Qiymat','Holat']
    ws.append(headers)
    for p in Part.objects.filter(company=c, **branch_kwargs(request)).order_by('category','name','code'):
        x=_part_json(p); ws.append([x['code'],x['name'],x['category'],x['model'],x['unit'],x['stock'],x['min_stock'],x['unit_price'],x['stock_value'],x['status']])
    for cell in ws[1]:
        cell.font=Font(bold=True); cell.alignment=Alignment(horizontal='center')
    response=HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition']='attachment; filename="avtopark_pro_sklad.xlsx"'
    wb.save(response)
    return response

@csrf_protect
@require_http_methods(['POST'])
def warehouse_return_create(request):
    if not request.user.is_authenticated or not _warehouse_direct(request):
        return JsonResponse({'error':'Ruxsat yo‘q'}, status=403)
    c=_company(request); data=_json(request) or {}
    try:
        issue=Issue.objects.filter(company=c, **branch_kwargs(request), pk=data.get('issue_id')).select_related('part','vehicle').first()
        if not issue: raise ValueError('Chiqim topilmadi')
        qty=_num(data.get('quantity'),'Miqdor',0)
        if qty<=0: raise ValueError('Miqdor 0 dan katta bo‘lishi kerak')
        already=float(issue.returns.aggregate(v=Sum('quantity'))['v'] or 0)
        if qty > float(issue.quantity)-already+1e-9:
            raise ValueError(f'Qaytarish limiti oshib ketdi. Qolgan: {float(issue.quantity)-already:g} {issue.part.unit}')
        with transaction.atomic():
            r=WarehouseReturn.objects.create(
                company=c,branch=getattr(request,"active_branch",None),date=_date(data.get('date')),issue=issue,part=issue.part,vehicle=issue.vehicle,
                quantity=qty,unit_price=float(issue.unit_price),document_no=str(data.get('document_no','')).strip(),
                reason=str(data.get('reason','')).strip(),note=str(data.get('note','')).strip())
            # Reduce the linked Finance warehouse expense; never below zero.
            if issue.expense_id:
                exp=OtherExpense.objects.select_for_update().filter(pk=issue.expense_id,company=c).first()
                if exp:
                    exp.amount=max(0,float(exp.amount)-float(r.total)); exp.save(update_fields=['amount'])
    except ValueError as e:
        return JsonResponse({'error':str(e)},status=400)
    _audit(request, c, 'CREATE', r, f'Ombor qaytarishi yaratildi: {r.part.name}', {'quantity': float(r.quantity), 'issue_id': r.issue_id})
    return JsonResponse({'return':_return_json(r)},status=201)

@csrf_protect
@require_http_methods(['POST'])
def warehouse_adjustment_create(request):
    if not request.user.is_authenticated or not _warehouse_direct(request):
        return JsonResponse({'error':'Ruxsat yo‘q'}, status=403)
    c=_company(request); data=_json(request) or {}
    try:
        part=_get_part(request,c,data.get('part_id'))
        if not part: raise ValueError('Mahsulot topilmadi')
        qty=_num(data.get('quantity'),'Korreksiya miqdori',None)
        if qty == 0: raise ValueError('Korreksiya 0 bo‘lishi mumkin emas')
        price=_num(data.get('unit_price',part.unit_price),'Birlik narxi',0)
        current=_stock_for_part(part)
        if current+qty < -1e-9: raise ValueError(f'Kamomad qoldiqdan oshib ketdi. Mavjud: {current:g} {part.unit}')
        a=WarehouseAdjustment.objects.create(
            company=c,branch=getattr(request,"active_branch",None),date=_date(data.get('date')),part=part,quantity=qty,unit_price=price,
            reason=str(data.get('reason','')).strip() or 'Inventarizatsiya',
            document_no=str(data.get('document_no','')).strip(),note=str(data.get('note','')).strip())
    except ValueError as e:
        return JsonResponse({'error':str(e)},status=400)
    _audit(request, c, 'CREATE', a, f'Ombor inventarizatsiya korreksiyasi: {a.part.name}', {'quantity': float(a.quantity), 'reason': a.reason})
    return JsonResponse({'adjustment':_adjustment_json(a)},status=201)

@require_http_methods(['GET'])
def warehouse_card(request, part_id):
    if not request.user.is_authenticated or not _can_manage(request):
        return JsonResponse({'error':'Ruxsat yo‘q'},status=403)
    c=_company(request); part=Part.objects.filter(company=c, **branch_kwargs(request),pk=part_id).first()
    if not part: return JsonResponse({'error':'Mahsulot topilmadi'},status=404)
    rows=[]
    for r in Receipt.objects.filter(company=c, **branch_kwargs(request),part=part).order_by('date','id'):
        rows.append({'date':r.date.isoformat(),'type':'KIRIM','document':r.document_no,'quantity':float(r.quantity),'unit_price':float(r.unit_price),'total':float(r.total),'counterparty':r.supplier})
    for i in Issue.objects.filter(company=c, **branch_kwargs(request),part=part).order_by('date','id'):
        rows.append({'date':i.date.isoformat(),'type':'CHIQIM','document':'','quantity':-float(i.quantity),'unit_price':float(i.unit_price),'total':-float(i.total),'counterparty':i.vehicle.number})
    for r in WarehouseReturn.objects.filter(company=c, **branch_kwargs(request),part=part).order_by('date','id'):
        rows.append({'date':r.date.isoformat(),'type':'QAYTARISH','document':r.document_no,'quantity':float(r.quantity),'unit_price':float(r.unit_price),'total':float(r.total),'counterparty':r.vehicle.number})
    for a in WarehouseAdjustment.objects.filter(company=c, **branch_kwargs(request),part=part).order_by('date','id'):
        rows.append({'date':a.date.isoformat(),'type':'KORREKSIYA','document':a.document_no,'quantity':float(a.quantity),'unit_price':float(a.unit_price),'total':float(a.quantity)*float(a.unit_price),'counterparty':a.reason})
    balance=0
    for x in sorted(rows,key=lambda x:(x['date'],x['type'])):
        balance += x['quantity']; x['balance']=balance
    return JsonResponse({'part':_part_json(part),'rows':rows})
