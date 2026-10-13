import csv
import json
import base64
import binascii
from django.core.files.base import ContentFile
from datetime import date
from django.utils import timezone
from django.core.files.storage import default_storage
from core.models import GasStation, DriverProfile, CompanyMembership
from django.db import transaction
from django.db.models import Sum
from django.http import JsonResponse, HttpResponse
from django.views.decorators.csrf import csrf_protect
from django.views.decorators.http import require_http_methods

from core.models import FuelEntry, Vehicle, UserProfile
from .trips import _company, _can_manage
from .branch_scope import branch_kwargs, relation_branch_kwargs


def _json(request):
    try:
        return json.loads(request.body or '{}')
    except json.JSONDecodeError:
        return None


def _num(value, name, minimum=None, required=True):
    if value in (None, ''):
        if required:
            raise ValueError(f'{name} majburiy')
        return None
    try:
        n = float(value)
    except (TypeError, ValueError):
        raise ValueError(f'{name} noto\'g\'ri')
    if minimum is not None and n < minimum:
        raise ValueError(f'{name} {minimum:g} dan kichik bo\'lishi mumkin emas')
    return n


def _date(value):
    if not value:
        return date.today()
    try:
        return date.fromisoformat(str(value))
    except (TypeError, ValueError):
        raise ValueError('Sana noto\'g\'ri')


def _vehicle(request, c, value):
    try:
        return Vehicle.objects.get(company=c, **branch_kwargs(request), pk=int(value), active=True)
    except (Vehicle.DoesNotExist, TypeError, ValueError):
        return None


def _previous_entry(request, c, vehicle, dt, pk=None):
    qs = FuelEntry.objects.filter(company=c, vehicle=vehicle).filter(date__lte=dt).order_by('-date', '-id')
    if pk:
        qs = qs.exclude(pk=pk)
    return qs.first()


def _row(c, x):
    previous = _previous_entry(None, c, x.vehicle, x.date, x.pk)
    distance = None
    consumption = None
    status = 'NORMAL'
    status_text = 'Normal'
    if previous:
        distance = float(x.odometer) - float(previous.odometer)
        if distance > 0:
            consumption = float(x.liters) * 100 / distance
            if consumption > 80:
                status, status_text = 'HIGH', 'Sarf juda yuqori'
            elif consumption < 3:
                status, status_text = 'LOW', 'Sarf juda past — tekshiring'
        elif distance == 0:
            status, status_text = 'NO_DISTANCE', 'Oldingi yoqilg\'idan beri masofa yo\'q'
        else:
            status, status_text = 'ODOMETER', 'Probeg oldingi yozuvdan kichik'
    return {
        'id': x.id, 'date': x.date.isoformat(), 'vehicle_id': x.vehicle_id,
        'vehicle': x.vehicle.number, 'model': x.vehicle.model,
        'odometer': float(x.odometer), 'liters': float(x.liters),
        'price_per_liter': float(x.price_per_liter), 'total': float(x.total),
        'station': x.station, 'fuel_type': x.fuel_type,
        'document_no': x.document_no, 'receipt_image': x.receipt_image.url if x.receipt_image else '', 'quantity_unit': x.quantity_unit, 'submitted_by': x.submitted_by.get_full_name() if x.submitted_by else '', 'submitted_at': timezone.localtime(x.submitted_at).isoformat() if x.submitted_at else '', 'gas_station_id': x.gas_station_id, 'ocr_text': x.ocr_text, 'ocr_quantity_m3': x.ocr_quantity_m3, 'ocr_status': x.ocr_status, 'distance_since_previous': distance,
        'consumption_l100': consumption, 'status': status, 'status_text': status_text,
        'previous_odometer': float(previous.odometer) if previous else None,
    }


def _response(request, c, vehicle_id=None, vehicle_ids=None, date_from=None, date_to=None, limit=5000):
    qs = FuelEntry.objects.filter(company=c, **branch_kwargs(request)).select_related('vehicle', 'submitted_by')
    if vehicle_ids:
        qs = qs.filter(vehicle_id__in=vehicle_ids)
    elif vehicle_id:
        qs = qs.filter(vehicle_id=vehicle_id)
    if date_from:
        qs = qs.filter(date__gte=date_from)
    if date_to:
        qs = qs.filter(date__lte=date_to)
    qs = qs.order_by('-date', '-id')
    rows = [_row(c, x) for x in qs[:limit]]
    total_liters = sum(x['liters'] for x in rows)
    total_cost = sum(x['total'] for x in rows)
    measured = [x['consumption_l100'] for x in rows if x['consumption_l100'] is not None and x['status'] not in ('ODOMETER', 'NO_DISTANCE')]
    avg_consumption = (sum(measured) / len(measured)) if measured else None
    return {
        'results': rows, 'count': len(rows),
        'summary': {'entries': len(rows), 'liters': total_liters, 'cost': total_cost,
                    'avg_consumption_l100': avg_consumption,
                    'alerts': sum(1 for x in rows if x['status'] != 'NORMAL')},
        'vehicles': [{'id': v.id, 'number': v.number, 'model': v.model,
                      'mileage': float(v.current_mileage or 0)}
                     for v in Vehicle.objects.filter(company=c, **branch_kwargs(request), active=True).order_by('number')],
        'driver_summary': _driver_summary(rows),
    }


def _driver_summary(rows):
    grouped = {}
    for row in rows:
        name = row.get('submitted_by') or 'Noma’lum foydalanuvchi'
        item = grouped.setdefault(name, {'driver': name, 'entries': 0, 'photos': 0, 'quantity_m3': 0.0, 'vehicles': set()})
        item['entries'] += 1
        item['photos'] += 1 if row.get('receipt_image') else 0
        if row.get('quantity_unit') == 'M3':
            item['quantity_m3'] += float(row.get('liters') or 0)
        item['vehicles'].add(row.get('vehicle') or '')
    return [{'driver': x['driver'], 'entries': x['entries'], 'photos': x['photos'],
             'quantity_m3': round(x['quantity_m3'], 2), 'vehicles_count': len(x['vehicles'])}
            for x in sorted(grouped.values(), key=lambda v: v['entries'], reverse=True)]


@require_http_methods(['GET'])
def fuel_list(request):
    if not request.user.is_authenticated or not _can_manage(request):
        return JsonResponse({'error': 'Ruxsat yo\'q'}, status=403)
    c = _company(request)
    if not c:
        return JsonResponse({'results': [], 'count': 0, 'summary': {}, 'vehicles': []})
    try:
        vehicle_id = int(request.GET.get('vehicle_id')) if request.GET.get('vehicle_id') else None
    except ValueError:
        vehicle_id = None
    try:
        date_from = _date(request.GET.get('from')) if request.GET.get('from') else None
        date_to = _date(request.GET.get('to')) if request.GET.get('to') else None
    except ValueError as e:
        return JsonResponse({'error': str(e)}, status=400)
    vehicle_ids = []
    for raw_id in request.GET.getlist('vehicle_ids'):
        try:
            vehicle_ids.append(int(raw_id))
        except (TypeError, ValueError):
            return JsonResponse({'error': 'Avtomobil tanlovi noto‘g‘ri'}, status=400)
    return JsonResponse(_response(request, c, vehicle_id, vehicle_ids or None, date_from, date_to))


@csrf_protect
@require_http_methods(['POST'])
def fuel_create(request):
    return _save(request)


@csrf_protect
@require_http_methods(['PATCH'])
def fuel_update(request, pk):
    return _save(request, pk)


def _save(request, pk=None):
    if not request.user.is_authenticated or not _can_manage(request):
        return JsonResponse({'error': 'Ruxsat yo\'q'}, status=403)
    c = _company(request)
    data = _json(request) or {}
    instance = FuelEntry.objects.filter(company=c, **branch_kwargs(request), pk=pk).first() if pk else None
    if pk and not instance:
        return JsonResponse({'error': 'Yoqilg\'i yozuvi topilmadi'}, status=404)
    try:
        vehicle = _vehicle(request, c, data.get('vehicle_id'))
        if not vehicle:
            raise ValueError('Faol avtomobilni tanlang')
        dt = _date(data.get('date'))
        odometer = _num(data.get('odometer'), 'Probeg', 0)
        liters = _num(data.get('liters'), 'Litr', 0)
        if liters <= 0:
            raise ValueError('Litr 0 dan katta bo\'lishi kerak')
        price = _num(data.get('price_per_liter'), '1 litr narxi', 0)
        with transaction.atomic():
            x = instance or FuelEntry(company=c)
            x.company = c
            x.vehicle = vehicle
            x.date = dt
            x.odometer = odometer
            x.liters = liters
            x.price_per_liter = price
            x.station = str(data.get('station', '')).strip()
            x.fuel_type = str(data.get('fuel_type', '')).strip()
            x.document_no = str(data.get('document_no', '')).strip()
            x.quantity_unit = str(data.get('quantity_unit', 'M3' if x.fuel_type.lower() == 'gaz' else 'L')).upper()
            if x.quantity_unit not in ('L', 'M3'):
                raise ValueError('Miqdor birligini tanlang')
            x.submitted_by = request.user
            image_data = data.get('receipt_image_data')
            if image_data:
                try:
                    header, encoded = image_data.split(',', 1)
                    if not header.startswith('data:image/'):
                        raise ValueError()
                    raw = base64.b64decode(encoded, validate=True)
                    if len(raw) > 8 * 1024 * 1024:
                        raise ValueError('Rasm 8 MB dan kichik bo‘lishi kerak')
                    ext = 'png' if 'image/png' in header else 'webp' if 'image/webp' in header else 'jpg'
                    x.receipt_image.save(f'fuel-{request.user_id}-{dt.isoformat()}.{ext}', ContentFile(raw), save=False)
                except (ValueError, binascii.Error):
                    raise ValueError('Chek rasmi noto‘g‘ri yoki 8 MB dan katta')
            x.save()
            # Keep the vehicle's operational odometer at least as current as fuel data.
            if float(vehicle.current_mileage or 0) < odometer:
                vehicle.current_mileage = odometer
                vehicle.save(update_fields=['current_mileage'])
    except ValueError as e:
        return JsonResponse({'error': str(e)}, status=400)
    return JsonResponse({'fuel': _row(c, x)})


@require_http_methods(['GET', 'POST'])
@csrf_protect
def fuel_stations(request):
    if not request.user.is_authenticated:
        return JsonResponse({'error': 'Avval tizimga kiring'}, status=401)
    c = _company(request)
    if not c:
        return JsonResponse({'error': 'Kompaniya topilmadi'}, status=403)
    if request.method == 'GET':
        stations = GasStation.objects.filter(company=c, active=True, **branch_kwargs(request)).order_by('name')
        return JsonResponse({'results': [{'id': s.id, 'name': s.name, 'address': s.address} for s in stations]})
    if not _can_manage(request):
        return JsonResponse({'error': 'Faqat rahbar stansiya qo‘sha oladi'}, status=403)
    data = _json(request) or {}
    name = str(data.get('name', '')).strip()
    if not name:
        return JsonResponse({'error': 'Zapravka nomi majburiy'}, status=400)
    station, created = GasStation.objects.get_or_create(company=c, name=name, defaults={'address': str(data.get('address', '')).strip(), 'branch': getattr(request, 'active_branch', None)})
    if not created and not station.active:
        station.active = True; station.save(update_fields=['active'])
    return JsonResponse({'station': {'id': station.id, 'name': station.name, 'address': station.address}, 'created': created}, status=201 if created else 200)


@csrf_protect
@require_http_methods(['POST'])
def driver_gas_submit(request):
    """Driver submits a dispenser photo; date/time are server-controlled."""
    if not request.user.is_authenticated:
        return JsonResponse({'error': 'Avval tizimga kiring'}, status=401)
    c = _company(request)
    profile = UserProfile.objects.filter(user=request.user).first()
    membership = CompanyMembership.objects.filter(user=request.user, company=c, status='ACTIVE').first() if c else None
    role = membership.role if membership else (profile.role if profile else None)
    if not c or role not in (UserProfile.ROLE_DRIVER, UserProfile.ROLE_DIRECTOR, UserProfile.ROLE_MANAGER):
        return JsonResponse({'error': 'Bu amal uchun ruxsat yo‘q'}, status=403)
    data = _json(request) or {}
    try:
        quantity = _num(data.get('quantity_m3'), 'Gaz miqdori (m³)', 0.01)
        station = GasStation.objects.get(company=c, active=True, pk=int(data.get('station_id')))
        driver = DriverProfile.objects.filter(company=c, user=request.user, active=True).select_related('vehicle').first()
        if role == UserProfile.ROLE_DRIVER:
            vehicle = driver.vehicle if driver else None
        else:
            vehicle = _vehicle(request, c, data.get('vehicle_id'))
        if not vehicle:
            raise ValueError('Sizga biriktirilgan faol avtomobil topilmadi')
        if vehicle.company_id != c.id:
            raise ValueError('Avtomobil ushbu kompaniyaga tegishli emas')
        image_data = data.get('dispenser_image_data')
        if not image_data:
            raise ValueError('Gaz nasosi displeyi rasmi majburiy')
        header, encoded = image_data.split(',', 1)
        if not header.startswith('data:image/'):
            raise ValueError('Faqat rasm fayli yuklash mumkin')
        raw = base64.b64decode(encoded, validate=True)
        if not raw or len(raw) > 8 * 1024 * 1024:
            raise ValueError('Rasm 8 MB dan kichik bo‘lishi kerak')
        ext = 'png' if 'image/png' in header else 'webp' if 'image/webp' in header else 'jpg'
        now = timezone.now()
        entry = FuelEntry(company=c, branch=getattr(request, 'active_branch', None), vehicle=vehicle,
            date=timezone.localdate(now), odometer=float(vehicle.current_mileage or 0), liters=quantity,
            price_per_liter=0, station=station.name, gas_station=station, fuel_type='Gaz', quantity_unit='M3',
            submitted_by=request.user, submitted_at=now, ocr_status='PENDING')
        entry.receipt_image.save(f'dispenser-{vehicle.id}-{now.strftime("%Y%m%d-%H%M%S")}.{ext}', ContentFile(raw), save=False)
        entry.save()
        # OCR is optional at runtime. Original photo and manually entered value are always retained.
        try:
            import pytesseract
            from PIL import Image, ImageOps, ImageEnhance
            import io, re
            img = Image.open(io.BytesIO(raw)).convert('L')
            img = ImageOps.autocontrast(img)
            img = ImageEnhance.Contrast(img).enhance(1.8)
            text = pytesseract.image_to_string(img, config='--psm 6')[:4000]
            nums = re.findall(r'(?<!\d)(\d{1,5}(?:[.,]\d{1,3})?)(?!\d)', text.replace(' ', ''))
            values = [float(v.replace(',', '.')) for v in nums]
            entry.ocr_text = text
            if values:
                decimal_values = [v for token, v in zip(nums, values) if '.' in token or ',' in token]
                # Prefer an explicitly decimal reading; otherwise use the final OCR token.
                candidate = decimal_values[-1] if decimal_values else values[-1]
                entry.ocr_quantity_m3 = candidate
                entry.ocr_status = 'MATCH' if abs(candidate - quantity) <= max(0.5, quantity * 0.03) else 'MISMATCH'
            else:
                entry.ocr_status = 'UNREADABLE'
            entry.save(update_fields=['ocr_text', 'ocr_quantity_m3', 'ocr_status'])
        except Exception:
            # Missing OCR binary or unreadable image must not discard a valid submission.
            entry.ocr_status = 'PENDING'; entry.save(update_fields=['ocr_status'])
        return JsonResponse({'fuel': _row(c, entry)}, status=201)
    except (ValueError, TypeError, GasStation.DoesNotExist, binascii.Error) as e:
        return JsonResponse({'error': str(e) or 'Ma’lumot noto‘g‘ri'}, status=400)


@require_http_methods(['GET'])
def fuel_export(request):
    if not request.user.is_authenticated or not _can_manage(request):
        return JsonResponse({'error': 'Ruxsat yo\'q'}, status=403)
    c = _company(request)
    if not c:
        return JsonResponse({'error': 'Kompaniya topilmadi'}, status=403)
    try:
        date_from = _date(request.GET.get('from')) if request.GET.get('from') else None
        date_to = _date(request.GET.get('to')) if request.GET.get('to') else None
        vehicle_ids = [int(x) for x in request.GET.getlist('vehicle_ids')]
        if any(not Vehicle.objects.filter(company=c, **branch_kwargs(request), pk=vid).exists() for vid in vehicle_ids):
            return JsonResponse({'error': 'Tanlangan avtomobil kompaniya yoki filialga tegishli emas'}, status=400)
    except ValueError as e:
        return JsonResponse({'error': str(e)}, status=400)
    data = _response(request, c, vehicle_ids=vehicle_ids or None, date_from=date_from, date_to=date_to, limit=50000)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        import io
        wb = Workbook()
        summary = wb.active; summary.title = 'Avtomobillar'
        summary.append(['Davlat raqami', 'Model', 'Zapravka soni', 'Rasm soni', 'Jami gaz (m³)', 'Haydovchilar soni'])
        groups = {}
        for x in data['results']:
            g = groups.setdefault(x['vehicle_id'], {'number': x['vehicle'], 'model': x['model'], 'entries': 0, 'photos': 0, 'm3': 0.0, 'drivers': set()})
            g['entries'] += 1; g['photos'] += bool(x['receipt_image'])
            if x['quantity_unit'] == 'M3': g['m3'] += x['liters']
            if x['submitted_by']: g['drivers'].add(x['submitted_by'])
        for g in sorted(groups.values(), key=lambda x: x['number']):
            summary.append([g['number'], g['model'], g['entries'], g['photos'], round(g['m3'], 2), len(g['drivers'])])
        details = wb.create_sheet('Zapravkalar')
        details.append(['Sana', 'Vaqt (Toshkent)', 'Davlat raqami', 'Model', 'Haydovchi', 'Zapravka', 'Miqdor', 'Birlik', 'Displey rasmi URL', 'OCR qiymati (m³)', 'OCR holati', 'OCR matni'])
        for x in data['results']:
            details.append([x['date'], x.get('submitted_at',''), x['vehicle'], x['model'], x['submitted_by'], x['station'], x['liters'], x['quantity_unit'], x['receipt_image'], x.get('ocr_quantity_m3'), x.get('ocr_status',''), x.get('ocr_text','')])
        for ws in wb.worksheets:
            ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='17365D')
                cell.alignment = Alignment(wrap_text=True, vertical='center')
            for col in ws.columns:
                width = min(max(max(len(str(c.value or '')) for c in col) + 2, 12), 42)
                ws.column_dimensions[get_column_letter(col[0].column)].width = width
        out = io.BytesIO(); wb.save(out); out.seek(0)
        response = HttpResponse(out.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="avtopark_yokilgi_hisoboti.xlsx"'
        return response
    except ImportError:
        return JsonResponse({'error': 'Excel eksporti uchun openpyxl kutubxonasini o‘rnating'}, status=500)
